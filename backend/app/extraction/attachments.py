import io, zipfile, subprocess, sys, json
from pathlib import Path

def extract_attachment(data,mime):
    if len(data)>10*1024*1024 or mime not in ('application/pdf','application/vnd.openxmlformats-officedocument.wordprocessingml.document'):
        return {'status':'manual_review','text':'','reason':'Unsupported attachment or larger than 10 MB'}
    try:
        if mime.endswith('document'):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(x.file_size for x in archive.infolist())>50*1024*1024:raise ValueError('Expanded document is too large')
        result=subprocess.run([sys.executable,str(Path(__file__).resolve()),mime],input=data,capture_output=True,timeout=30)
        if result.returncode:raise ValueError('Unable to read attachment')
        return json.loads(result.stdout)
    except Exception:return {'status':'manual_review','text':'','reason':'Attachment needs manual review'}

def parse(data,mime):
    truncated=False
    if mime=='application/pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data));truncated=len(reader.pages)>50
        text='\n'.join(p.extract_text() or '' for p in reader.pages[:50])
    else:
        from docx import Document
        doc=Document(io.BytesIO(data));text='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    truncated=truncated or len(text)>100000
    return {'status':'manual_review' if truncated or not text.strip() else 'parsed','text':text[:100000],'reason':'Truncated or scanned document' if truncated or not text.strip() else ''}
if __name__=='__main__':
    # Production runs on Linux. Bound the parser child, not the API/worker parent,
    # so a compressed or hostile document becomes manual review rather than OOM.
    if sys.platform!='win32':
        import resource
        resource.setrlimit(resource.RLIMIT_AS,(256*1024*1024,256*1024*1024))
        resource.setrlimit(resource.RLIMIT_CPU,(20,25))
    print(json.dumps(parse(sys.stdin.buffer.read(),sys.argv[1])))
